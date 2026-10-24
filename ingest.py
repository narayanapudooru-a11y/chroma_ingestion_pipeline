"""
Incremental local-file -> ChromaDB ingestion pipeline.

Supports PDF and DOCX files. Runs once per invocation:
- Adds new files
- Re-ingests changed files
- Removes vectors for files deleted from the source directory
- Uses deterministic chunk IDs and a persistent Chroma collection

Recommended scheduling: invoke this script every Saturday at 08:00 Asia/Kolkata
using cron/systemd timer or Windows Task Scheduler. Do not keep a scheduler loop
inside the ingestion process unless you specifically need one.
"""
from __future__ import annotations

import hashlib
import logging
import os
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import chromadb
from chromadb.config import Settings
from pypdf import PdfReader
from docx import Document

# ---------------------------- Configuration ----------------------------

SOURCE_DIR = Path(os.getenv("INGEST_SOURCE_DIR", r"D:\knowledge_documents")).expanduser().resolve()
CHROMA_DIR = Path(os.getenv("CHROMA_PERSIST_DIR", "./chroma_data")).expanduser().resolve()
COLLECTION_NAME = os.getenv("CHROMA_COLLECTION", "local_documents")
EMBEDDING_MODEL = os.getenv("EMBEDDING_MODEL", "sentence-transformers/all-MiniLM-L6-v2")

# Tune for your documents / embedding model.
CHUNK_SIZE_CHARS = int(os.getenv("CHUNK_SIZE_CHARS", "1800"))
CHUNK_OVERLAP_CHARS = int(os.getenv("CHUNK_OVERLAP_CHARS", "250"))
EMBED_BATCH_SIZE = int(os.getenv("EMBED_BATCH_SIZE", "64"))
MAX_RETRIES = int(os.getenv("MAX_RETRIES", "3"))
RETRY_DELAY_SECONDS = float(os.getenv("RETRY_DELAY_SECONDS", "2"))
SUPPORTED_EXTENSIONS = {".pdf", ".docx"}

LOG_DIR = Path(os.getenv("INGEST_LOG_DIR", "./logs")).expanduser().resolve()
LOG_DIR.mkdir(parents=True, exist_ok=True)
logging.basicConfig(
    level=os.getenv("LOG_LEVEL", "INFO").upper(),
    format="%(asctime)s %(levelname)s %(message)s",
    handlers=[
        logging.StreamHandler(sys.stdout),
        logging.FileHandler(LOG_DIR / "ingestion.log", encoding="utf-8"),
    ],
)
log = logging.getLogger("chroma_ingestion")


@dataclass
class SourceFile:
    path: Path
    relative_path: str
    sha256: str
    mtime: float
    size: int


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def discover_files(root: Path) -> dict[str, SourceFile]:
    if not root.exists() or not root.is_dir():
        raise RuntimeError(f"Source directory does not exist or is not a directory: {root}")

    found: dict[str, SourceFile] = {}
    for path in root.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue
        # Ignore common temporary Office files.
        if path.name.startswith("~$"):
            continue
        try:
            stat = path.stat()
            relative = path.relative_to(root).as_posix()
            found[relative] = SourceFile(
                path=path,
                relative_path=relative,
                sha256=sha256_file(path),
                mtime=stat.st_mtime,
                size=stat.st_size,
            )
        except (OSError, PermissionError) as exc:
            # Fail the whole run rather than deleting vectors for an unreadable file.
            raise RuntimeError(f"Cannot scan/read source file {path}: {exc}") from exc
    return found


def extract_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    pages = []
    for page_no, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            log.warning("Could not extract page %s from %s: %s", page_no, path, exc)
            text = ""
        if text.strip():
            pages.append(f"[Page {page_no}]\n{text}")
    return "\n\n".join(pages)


def extract_docx(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    # Include table cell text, which is commonly missed by paragraph-only extraction.
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix == ".docx":
        return extract_docx(path)
    raise ValueError(f"Unsupported extension: {suffix}")


def chunk_text(text: str, size: int, overlap: int) -> list[str]:
    text = "\n".join(line.rstrip() for line in text.replace("\x00", "").splitlines()).strip()
    if not text:
        return []
    if size <= 0 or overlap < 0 or overlap >= size:
        raise ValueError("CHUNK_SIZE_CHARS must be > 0 and overlap must be >= 0 and < chunk size")

    chunks = []
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        # Prefer a paragraph/newline or sentence boundary near the end.
        if end < len(text):
            boundary = max(text.rfind("\n", start + size // 2, end),
                           text.rfind(". ", start + size // 2, end))
            if boundary > start:
                end = boundary + 1
        piece = text[start:end].strip()
        if piece:
            chunks.append(piece)
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return chunks


def stable_chunk_id(relative_path: str, content_hash: str, index: int) -> str:
    raw = f"{relative_path}\0{content_hash}\0{index}".encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


def retry(operation, description: str):
    last_error = None
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            return operation()
        except Exception as exc:
            last_error = exc
            if attempt < MAX_RETRIES:
                wait = RETRY_DELAY_SECONDS * attempt
                log.warning("%s failed (attempt %s/%s): %s; retrying in %ss",
                            description, attempt, MAX_RETRIES, exc, wait)
                time.sleep(wait)
    raise RuntimeError(f"{description} failed after {MAX_RETRIES} attempts") from last_error


def get_all_existing_metadata(collection) -> list[dict]:
    """Fetch metadata only, in pages, to identify files represented in Chroma."""
    result = []
    offset = 0
    page_size = 1000
    while True:
        batch = collection.get(
            limit=page_size,
            offset=offset,
            include=["metadatas"],
        )
        ids = batch.get("ids", [])
        if not ids:
            break
        metadatas = batch.get("metadatas") or []
        result.extend(m for m in metadatas if m)
        offset += len(ids)
        if len(ids) < page_size:
            break
    return result


def delete_source_chunks(collection, relative_path: str) -> int:
    existing = collection.get(
        where={"source_path": relative_path},
        include=[],
    )
    ids = existing.get("ids", [])
    if ids:
        retry(lambda: collection.delete(ids=ids), f"delete vectors for {relative_path}")
    return len(ids)


def upsert_document(collection, model, source: SourceFile) -> int:
    text = retry(lambda: extract_text(source.path), f"extract {source.relative_path}")
    chunks = chunk_text(text, CHUNK_SIZE_CHARS, CHUNK_OVERLAP_CHARS)
    if not chunks:
        # Empty/scanned PDFs can be image-only. Avoid silently removing an older
        # successful version: fail this file and keep the previous vectors.
        raise RuntimeError(
            f"No extractable text in {source.relative_path}. It may be scanned/image-only; "
            "run OCR before ingestion. Existing vectors were not intentionally replaced."
        )

    ids = [stable_chunk_id(source.relative_path, source.sha256, i) for i in range(len(chunks))]
    metadatas = [{
        "source_path": source.relative_path,
        "source_name": source.path.name,
        "file_hash": source.sha256,
        "chunk_index": i,
        "file_mtime": source.mtime,
        "file_size": source.size,
        "file_type": source.path.suffix.lower().lstrip("."),
    } for i in range(len(chunks))]

    # Embed and upsert new version first. If embedding/upsert fails, old vectors remain.
    for start in range(0, len(chunks), EMBED_BATCH_SIZE):
        batch_chunks = chunks[start:start + EMBED_BATCH_SIZE]
        batch_ids = ids[start:start + EMBED_BATCH_SIZE]
        batch_meta = metadatas[start:start + EMBED_BATCH_SIZE]
        embeddings = retry(
            lambda: model.encode(batch_chunks, normalize_embeddings=True).tolist(),
            f"embed {source.relative_path} batch {start}",
        )
        retry(
            lambda: collection.upsert(
                ids=batch_ids,
                documents=batch_chunks,
                metadatas=batch_meta,
                embeddings=embeddings,
            ),
            f"upsert {source.relative_path} batch {start}",
        )

    # Remove any previous version's chunks, preserving the newly upserted IDs.
    old = collection.get(where={"source_path": source.relative_path}, include=["metadatas"])
    old_ids = old.get("ids", [])
    old_meta = old.get("metadatas") or []
    stale_ids = [
        item_id for item_id, meta in zip(old_ids, old_meta)
        if meta and meta.get("file_hash") != source.sha256
    ]
    if stale_ids:
        retry(lambda: collection.delete(ids=stale_ids), f"remove stale chunks for {source.relative_path}")
    return len(chunks)


def run_ingestion() -> None:
    if CHUNK_OVERLAP_CHARS >= CHUNK_SIZE_CHARS:
        raise ValueError("CHUNK_OVERLAP_CHARS must be smaller than CHUNK_SIZE_CHARS")

    log.info("Starting ingestion. Source=%s Chroma=%s Collection=%s", SOURCE_DIR, CHROMA_DIR, COLLECTION_NAME)
    sources = discover_files(SOURCE_DIR)
    log.info("Discovered %d supported files", len(sources))

    CHROMA_DIR.mkdir(parents=True, exist_ok=True)
    client = chromadb.PersistentClient(
        path=str(CHROMA_DIR),
        settings=Settings(anonymized_telemetry=False),
    )
    collection = client.get_or_create_collection(
        name=COLLECTION_NAME,
        metadata={"hnsw:space": "cosine", "description": "Incremental local PDF/DOCX ingestion"},
    )
    from sentence_transformers import SentenceTransformer

    model = SentenceTransformer(EMBEDDING_MODEL)

    # Build current state by source_path and file_hash from Chroma metadata.
    existing_meta = get_all_existing_metadata(collection)
    existing_hashes: dict[str, set[str]] = {}
    for meta in existing_meta:
        source_path = meta.get("source_path")
        file_hash = meta.get("file_hash")
        if source_path and file_hash:
            existing_hashes.setdefault(source_path, set()).add(file_hash)

    added = changed = unchanged = failed = deleted_files = deleted_chunks = 0

    for relative_path, source in sorted(sources.items()):
        hashes = existing_hashes.get(relative_path, set())
        if hashes == {source.sha256}:
            unchanged += 1
            continue
        is_new = not hashes
        try:
            count = upsert_document(collection, model, source)
            if is_new:
                added += 1
                log.info("ADDED %s (%d chunks)", relative_path, count)
            else:
                changed += 1
                log.info("UPDATED %s (%d chunks)", relative_path, count)
        except Exception:
            failed += 1
            log.exception("FAILED %s; continuing with other files", relative_path)

    # Delete only paths that are absent from the successful directory scan.
    current_paths = set(sources.keys())
    existing_paths = set(existing_hashes.keys())
    for missing_path in sorted(existing_paths - current_paths):
        try:
            count = delete_source_chunks(collection, missing_path)
            deleted_files += 1
            deleted_chunks += count
            log.info("DELETED %s (%d chunks)", missing_path, count)
        except Exception:
            failed += 1
            log.exception("FAILED to delete vectors for missing file %s", missing_path)

    log.info(
        "Finished. added=%d changed=%d unchanged=%d deleted_files=%d deleted_chunks=%d failed=%d total_vectors=%d",
        added, changed, unchanged, deleted_files, deleted_chunks, failed, collection.count(),
    )
    if failed:
        # Non-zero exit lets Task Scheduler / monitoring detect a partial failure.
        raise RuntimeError(f"Ingestion completed with {failed} failure(s); check the log.")


if __name__ == "__main__":
    try:
        run_ingestion()
    except Exception:
        log.exception("Ingestion run failed")
        sys.exit(1)
